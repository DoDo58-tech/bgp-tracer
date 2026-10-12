#!/usr/bin/env python3
"""Regenerate the CSS + HCAF method overview (fig2) and export to PPTX.
"""
from __future__ import annotations

import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

OUT_DIR = ROOT / "docs" / "paper" / "figs"
OUT_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "#0b2545"
SKY = "#1d6fb8"
TEAL = "#0e8a73"
AMBER = "#e08520"
SLATE = "#5a6a7a"
LIGHT = "#e8eff7"
WARM = "#fde9d2"
COOL = "#e0f0fa"

plt.rcParams.update({
    "font.size": 10, "font.family": "DejaVu Sans",
    "figure.dpi": 200, "savefig.dpi": 300, "pdf.fonttype": 42,
})


def save(fig, name):
    fig.savefig(OUT_DIR / name, format="pdf", bbox_inches="tight", pad_inches=0.15)
    plt.close(fig)
    print(f"Wrote {name}")


# ── PDF (matplotlib) ───────────────────────────────────────────────


def fig2_pdf():
    fig, ax = plt.subplots(figsize=(7.6, 5.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 10)
    ax.axis("off")
    ax.set_facecolor("#ffffff")

    # Title bar
    ax.add_patch(Rectangle((0, 9.30), 10, 0.7, facecolor=NAVY))
    ax.text(5, 9.65, "Causal Attribution Pipeline",
            ha="center", va="center", fontsize=13, fontweight="bold",
            color="white")
    ax.text(5, 8.95, "Pearl-PNS estimator  +  4-channel Bayesian synthesis",
            ha="center", fontsize=9.5, color=SLATE, style="italic")

    # ── INPUT ──
    ax.text(1.5, 8.55, "INPUT", ha="center", fontsize=9, color=SKY,
            fontweight="bold")
    inp = FancyBboxPatch((0.20, 6.7), 2.6, 1.6,
                         boxstyle="round,pad=0.04,rounding_size=0.12",
                         linewidth=1.8, edgecolor=SKY, facecolor="white")
    ax.add_patch(inp)
    ax.text(1.5, 7.75, "Aligned timelines", ha="center", fontsize=11,
            fontweight="bold", color=SKY)
    ax.text(1.5, 7.30, r"$R(t)$ , $T(t)$", ha="center", fontsize=11)
    ax.text(1.5, 6.95, "5-min Radar buckets", ha="center", fontsize=8.5,
            color=SLATE, style="italic")

    # ── ESTIMATOR (Pearl-PNS) ──
    est = FancyBboxPatch((3.30, 5.8), 4.0, 2.5,
                         boxstyle="round,pad=0.04,rounding_size=0.12",
                         linewidth=2.0, edgecolor=NAVY, facecolor=COOL)
    ax.add_patch(est)
    ax.text(5.30, 7.95, "Pearl-PNS Estimator", ha="center", fontsize=12,
            fontweight="bold", color=NAVY)
    ax.text(5.30, 7.55, "(non-parametric, pre/post windows)",
            ha="center", fontsize=8.5, color=SLATE, style="italic")
    # Two sub-estimates side by side
    sub1 = FancyBboxPatch((3.50, 6.20), 1.75, 1.05,
                          boxstyle="round,pad=0.04,rounding_size=0.10",
                          linewidth=1.4, edgecolor=NAVY, facecolor="white")
    ax.add_patch(sub1)
    ax.text(4.37, 6.95, r"$\widehat{\mathrm{PDM}}(R\!\to\!T)$",
            ha="center", fontsize=10, fontweight="bold", color=NAVY)
    ax.text(4.37, 6.55, "$\\hat\\Delta$ + bootstrap CI",
            ha="center", fontsize=8, color=SLATE)
    sub2 = FancyBboxPatch((5.35, 6.20), 1.75, 1.05,
                          boxstyle="round,pad=0.04,rounding_size=0.10",
                          linewidth=1.4, edgecolor=NAVY, facecolor="white")
    ax.add_patch(sub2)
    ax.text(6.22, 6.95, r"$\widehat{\mathrm{PDM}}(T\!\to\!R)$",
            ha="center", fontsize=10, fontweight="bold", color=NAVY)
    ax.text(6.22, 6.55, "$\\hat\\Delta$ + bootstrap CI",
            ha="center", fontsize=8, color=SLATE)
    ax.text(5.30, 5.95, "permutation $p$-value", ha="center", fontsize=8.5,
            color=SLATE, style="italic")

    # ── 4 EVIDENCE CHANNELS ──
    ax.text(8.7, 8.55, "EVIDENCE CHANNELS", ha="center", fontsize=9,
            color=AMBER, fontweight="bold")
    ch_box = FancyBboxPatch((7.40, 4.3), 2.5, 4.0,
                            boxstyle="round,pad=0.04,rounding_size=0.12",
                            linewidth=1.8, edgecolor=AMBER, facecolor=WARM)
    ax.add_patch(ch_box)
    channels = [
        ("MSBD", "Multi-scale breakpoint"),
        ("CTPR", "Causal transmission path"),
        ("CIC",  "Cascade impact count"),
        ("TE",   "Transfer entropy"),
    ]
    for i, (name, desc) in enumerate(channels):
        y = 7.6 - i * 0.75
        sub = FancyBboxPatch((7.60, y), 2.1, 0.60,
                             boxstyle="round,pad=0.02,rounding_size=0.08",
                             linewidth=1.0, edgecolor=AMBER, facecolor="white")
        ax.add_patch(sub)
        ax.text(7.75, y + 0.30, name, fontsize=10, fontweight="bold",
                color=AMBER, ha="left", va="center")
        ax.text(8.45, y + 0.30, desc, fontsize=8.5, color=SLATE,
                ha="left", va="center")

    # ── BAYESIAN SYNTHESIS ──
    b = FancyBboxPatch((0.3, 2.7), 9.4, 1.5,
                       boxstyle="round,pad=0.04,rounding_size=0.12",
                       linewidth=2.0, edgecolor=TEAL, facecolor="white")
    ax.add_patch(b)
    ax.text(5.0, 3.85, "Bayesian Synthesis  (Pearl-PNS anchored)",
            ha="center", fontsize=12, fontweight="bold", color=TEAL)
    ax.text(5.0, 3.40,
            r"$\Pr(\text{cause} \mid \text{evidence}) \;\propto\; \Pr(\text{cause}) \cdot \prod_{\text{ch}} \mathrm{LR}(\text{ch})$",
            ha="center", fontsize=12, color=NAVY)
    ax.text(5.0, 2.95,
            "prior  anchored on PDM direction  •  disagreeing channels get 0.4× discount",
            ha="center", fontsize=9, color=SLATE, style="italic")

    # ── VERDICT ──
    v = FancyBboxPatch((0.3, 1.0), 9.4, 1.3,
                       boxstyle="round,pad=0.04,rounding_size=0.12",
                       linewidth=2.0, edgecolor=TEAL, facecolor=LIGHT)
    ax.add_patch(v)
    ax.text(5.0, 1.85, "Final Verdict", ha="center", fontsize=13,
            fontweight="bold", color=TEAL)
    ax.text(5.0, 1.45,
            "argmax posterior cause  +  95% CI  +  necessity  +  sufficiency  +  permutation $p$",
            ha="center", fontsize=9.5, color=NAVY)

    # ── ARROWS ──
    def arr(x1, y1, x2, y2, color):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                     mutation_scale=22, color=color, linewidth=2.2))
    arr(2.80, 7.5, 3.30, 7.5, color=SKY)         # input -> estimator
    arr(5.30, 5.8, 5.30, 4.2, color=NAVY)        # estimator -> bayes
    arr(8.70, 4.3, 7.0, 3.7, color=AMBER)        # channels -> bayes
    arr(5.0, 2.7, 5.0, 2.3, color=TEAL)          # bayes -> verdict

    # Legend strip
    ax.add_patch(Rectangle((0, 0.10), 10, 0.45, facecolor=NAVY))
    items = [(1.0, "Input", SKY), (3.0, "Pearl-PNS", NAVY),
             (5.5, "Channels", AMBER), (7.5, "Synthesis", TEAL)]
    for x, lbl, c in items:
        ax.add_patch(Rectangle((x, 0.25), 0.20, 0.15, color=c))
        ax.text(x + 0.30, 0.32, lbl, ha="left", va="center", fontsize=9.5,
                color="white", fontweight="bold")

    save(fig, "fig2_method_overview.pdf")


# ── PPTX (editable) ─────────────────────────────────────────────────


def _box(slide, x, y, w, h, text, fill, line, font_size=10,
         bold=True, border_w=1.25, color=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    s.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
    s.line.width = Pt(border_w)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    for i, line_text in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line_text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(
            (color or line).lstrip("#"))
    return s


def _text(slide, x, y, w, h, text, size, color, bold=True, align="center"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT,
                   "center": PP_ALIGN.CENTER,
                   "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return tb


def _rect(slide, x, y, w, h, fill, line=None, border_w=0.5):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    if line:
        s.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
        s.line.width = Pt(border_w)
    else:
        s.line.fill.background()
    return s


def _arrow(slide, x1, y1, x2, y2, color, weight=2.0):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    line.line.color.rgb = RGBColor.from_string(color.lstrip("#"))
    line.line.width = Pt(weight)
    from pptx.oxml.ns import qn
    from lxml import etree
    ln = line.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def fig2_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _rect(slide, 0, 0, 13.333, 7.5, "#ffffff", line=None)

    # Title
    _rect(slide, 0, 0, 13.333, 0.75, NAVY, line=None)
    _text(slide, 0, 0, 13.333, 0.75,
          "Causal Attribution Pipeline",
          size=22, color="#ffffff", bold=True)
    _text(slide, 0, 0.80, 13.333, 0.35,
          "Pearl-PNS estimator  +  4-channel Bayesian synthesis",
          size=12, color=SLATE, bold=False)

    # INPUT
    _text(slide, 0.2, 1.20, 2.6, 0.35, "INPUT",
          size=11, color=SKY, bold=True)
    _box(slide, 0.2, 1.60, 2.6, 1.6,
         "Aligned timelines\n$R(t)$ , $T(t)$\n5-min Radar buckets",
         fill="#ffffff", line=SKY, font_size=12, border_w=1.8,
         color=SKY)

    # ESTIMATOR
    _rect(slide, 3.2, 1.20, 4.0, 2.0, COOL, line=NAVY, border_w=2.0)
    _text(slide, 3.2, 1.30, 4.0, 0.4,
          "Pearl-PNS Estimator",
          size=14, color=NAVY, bold=True)
    _text(slide, 3.2, 1.70, 4.0, 0.30,
          "non-parametric, pre/post windows",
          size=10, color=SLATE, bold=False)
    _box(slide, 3.35, 2.10, 1.85, 0.95,
         r"PDM(R→T)" + "\n" + "$\\hat\\Delta$ + CI",
         fill="#ffffff", line=NAVY, font_size=10, border_w=1.2,
         color=NAVY)
    _box(slide, 5.25, 2.10, 1.85, 0.95,
         r"PDM(T→R)" + "\n" + "$\\hat\\Delta$ + CI",
         fill="#ffffff", line=NAVY, font_size=10, border_w=1.2,
         color=NAVY)

    # 4 CHANNELS
    _text(slide, 7.4, 1.20, 5.0, 0.35, "EVIDENCE CHANNELS",
          size=11, color=AMBER, bold=True)
    _rect(slide, 7.4, 1.60, 5.0, 1.6, WARM, line=AMBER, border_w=1.8)
    ch = [
        ("MSBD", "Multi-scale breakpoint"),
        ("CTPR", "Causal transmission path"),
    ]
    for i, (n, d) in enumerate(ch):
        x = 7.55 + i * 2.4
        _box(slide, x, 1.75, 2.3, 1.30, f"{n}\n{d}",
             fill="#ffffff", line=AMBER, font_size=10, border_w=1.0,
             color=AMBER)
    ch2 = [
        ("CIC",  "Cascade impact count"),
        ("TE",   "Transfer entropy"),
    ]
    for i, (n, d) in enumerate(ch2):
        x = 7.55 + i * 2.4
        _box(slide, x, 3.15, 2.3, 0.85, f"{n}: {d}",
             fill="#ffffff", line=AMBER, font_size=10, border_w=1.0,
             color=AMBER)

    # BAYESIAN SYNTHESIS
    _box(slide, 0.2, 3.30, 7.0, 1.5,
         "Bayesian Synthesis   (Pearl-PNS anchored)\n\n"
         "P(cause | evidence)  ∝  P(cause)  ·  ∏ LR(channel)",
         fill="#ffffff", line=TEAL, font_size=11, border_w=2.0,
         color=TEAL)
    _text(slide, 0.2, 4.55, 7.0, 0.30,
          "prior anchored on PDM direction  •  disagreeing channels × 0.4",
          size=10, color=SLATE, bold=False)

    # VERDICT
    _box(slide, 0.2, 5.10, 12.0, 1.3,
         "Final Verdict\n"
         "argmax posterior cause  +  95% CI  +  necessity  +  sufficiency  +  permutation p",
         fill=LIGHT, line=TEAL, font_size=12, border_w=2.0, color=TEAL)

    # ARROWS
    _arrow(slide, 2.80, 2.40, 3.20, 2.40, color=SKY, weight=2.0)
    _arrow(slide, 5.20, 3.20, 5.20, 3.30, color=NAVY, weight=2.0)
    _arrow(slide, 9.90, 3.20, 7.00, 3.60, color=AMBER, weight=2.0)
    _arrow(slide, 3.70, 4.80, 3.70, 5.10, color=TEAL, weight=2.0)
    _arrow(slide, 8.50, 4.80, 8.50, 5.10, color=TEAL, weight=2.0)

    # Legend
    _rect(slide, 0, 6.55, 13.333, 0.95, NAVY, line=None)
    items = [("Input", SKY), ("Pearl-PNS", NAVY),
             ("Channels", AMBER), ("Synthesis", TEAL)]
    for i, (lbl, c) in enumerate(items):
        x = 0.5 + i * 3.2
        _rect(slide, x, 6.85, 0.3, 0.30, c, line=None)
        _text(slide, x + 0.45, 6.78, 2.5, 0.45, lbl,
              size=12, color="#ffffff", bold=True, align="left")

    out = OUT_DIR / "fig2_method_overview.pptx"
    prs.save(out)
    print(f"Wrote {out.name}")


if __name__ == "__main__":
    fig2_pdf()
    fig2_pptx()
    print("Done.")
